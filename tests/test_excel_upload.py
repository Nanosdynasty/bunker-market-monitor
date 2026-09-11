from io import BytesIO

from openpyxl import Workbook

from bunker_market.excel_import import parse_workbook
from bunker_market.models import Observation, UploadState, db


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(["Port", "Fuel grade", "Price", "Timestamp", "Country"])
    sheet.append(["Singapore", "VLSFO", 625.5, "2026-09-11T08:30:00Z", "Singapore"])
    sheet.append(["Singapore", "IFO380", 482, "2026-09-11T08:30:00Z", "Singapore"])
    sheet.append(["Rotterdam", "MGO", 791, "2026-09-11T08:30:00Z", "Netherlands"])
    sheet["F2"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def csrf(client):
    client.get("/")
    with client.session_transaction() as session:
        return session["csrf_token"]


def test_parse_long_form_workbook(tmp_path):
    path = tmp_path / "prices.xlsx"
    path.write_bytes(workbook_bytes().getvalue())
    result = parse_workbook(path, "prices.xlsx")
    assert result.sheet_name == "Prices"
    assert result.layout == "long"
    assert len(result.observations) == 3
    assert {item.grade for item in result.observations} == {"VLSFO", "HSFO", "MGO"}
    assert result.formula_cache_missing == 1


def test_upload_imports_excel_as_a_source(app, client):
    response = client.post(
        "/api/uploads/excel",
        data={"file": (workbook_bytes(), "bunker prices.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": csrf(client)},
    )
    assert response.status_code == 200
    assert response.json["observations"] == 3
    with app.app_context():
        assert Observation.query.filter_by(provider_id="excel").count() == 3
        assert db.session.get(UploadState, 1).file_name == "bunker_prices.xlsx"
    dashboard = client.get("/api/dashboard").json
    assert dashboard["mode"] == "live"
    assert any(provider["id"] == "excel" for provider in dashboard["providers"])
    assert client.get("/api/sources").json["upload"]["worksheet"] == "Prices"


def test_upload_validation_and_csrf(client):
    assert client.post(
        "/api/uploads/excel",
        data={"file": (BytesIO(b"not excel"), "prices.xlsx")},
        content_type="multipart/form-data",
    ).status_code == 403
    response = client.post(
        "/api/uploads/excel",
        data={"file": (BytesIO(b"csv"), "prices.csv")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": csrf(client)},
    )
    assert response.status_code == 415


def test_failed_replacement_keeps_previous_excel_data(app, client):
    token = csrf(client)
    good = client.post(
        "/api/uploads/excel",
        data={"file": (workbook_bytes(), "good.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    assert good.status_code == 200
    bad = client.post(
        "/api/uploads/excel",
        data={"file": (BytesIO(b"broken"), "broken.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    assert bad.status_code == 422
    with app.app_context():
        assert Observation.query.filter_by(provider_id="excel").count() == 3
        assert db.session.get(UploadState, 1).status == "failed"


def test_parse_merged_heading_matrix(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Port matrix"
    sheet["B1"] = "Antwerp"
    sheet.merge_cells("B1:D1")
    sheet["B2"], sheet["C2"], sheet["D2"] = "HSFO", "VLSFO", "MGO"
    sheet.append(["2026-09-10", 470, 528, 672])
    path = tmp_path / "matrix.xlsx"
    workbook.save(path)
    result = parse_workbook(path, "matrix.xlsx")
    assert result.layout == "matrix"
    assert len(result.observations) == 3
    assert {item.port_name for item in result.observations} == {"Antwerp"}
