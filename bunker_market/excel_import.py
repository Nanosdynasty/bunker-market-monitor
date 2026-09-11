from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import Observation, Port, ProviderState, UploadState, db, utcnow
from .normalization import canonical_grade, parse_time, port_identity
from .providers.base import ProviderObservation


HEADER_ALIASES = {
    "port": {"port", "port name", "location", "market"},
    "grade": {"grade", "fuel", "fuel grade", "fuel type", "product"},
    "price": {"price", "rate", "price usd mt", "usd mt", "usd/mt"},
    "timestamp": {"timestamp", "date", "as of", "updated", "last updated", "time"},
    "country": {"country", "country name"},
    "region": {"region", "area"},
}


@dataclass
class ParseResult:
    file_name: str
    sheet_name: str
    layout: str
    observations: list[ProviderObservation]
    skipped_cells: int
    formula_cache_missing: int
    excel_errors: int
    sheet_names: list[str]


def _text(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _as_time(value: Any, fallback: datetime) -> datetime:
    if isinstance(value, datetime):
        return parse_time(value)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    return parse_time(value) if value not in (None, "") else fallback


def _formula_diagnostics(values_sheet, formulas_sheet):
    missing = 0
    errors = 0
    max_row = min(values_sheet.max_row, 10000)
    max_column = min(values_sheet.max_column, 200)
    value_rows = values_sheet.iter_rows(max_row=max_row, max_col=max_column)
    formula_rows = formulas_sheet.iter_rows(max_row=max_row, max_col=max_column)
    for value_row, formula_row in zip(value_rows, formula_rows):
        for value_cell, formula_cell in zip(value_row, formula_row):
            if formula_cell.data_type == "f" and value_cell.value is None:
                missing += 1
            if value_cell.data_type == "e" or (isinstance(value_cell.value, str) and value_cell.value.startswith("#")):
                errors += 1
    return missing, errors


def _looks_like_date(value: Any) -> bool:
    if isinstance(value, (date, datetime)):
        return True
    if not isinstance(value, str):
        return False
    text = value.strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d %b %Y", "%d-%b-%Y"):
        try:
            datetime.strptime(text, pattern)
            return True
        except ValueError:
            pass
    try:
        datetime.fromisoformat(text.replace("Z", "+00:00"))
        return True
    except ValueError:
        return False


def _find_columns(sheet):
    for row_number in range(1, min(sheet.max_row, 20) + 1):
        found = {}
        for column, cell in enumerate(sheet[row_number], 1):
            value = _text(cell.value)
            for field, aliases in HEADER_ALIASES.items():
                if value in aliases:
                    found[field] = column
        if {"port", "price"}.issubset(found) and "grade" in found:
            return row_number, found, "long"
        grade_columns = {
            column: canonical_grade(cell.value)
            for column, cell in enumerate(sheet[row_number], 1)
            if canonical_grade(cell.value)
        }
        if "port" in found and grade_columns:
            return row_number, {"port": found["port"], "grade_columns": grade_columns}, "wide"
    return None


def _parse_long(sheet, header_row, columns, fallback):
    rows = []
    skipped = 0
    for row_number in range(header_row + 1, min(sheet.max_row, 10000) + 1):
        port = sheet.cell(row_number, columns["port"]).value
        grade = canonical_grade(sheet.cell(row_number, columns["grade"]).value)
        raw_price = sheet.cell(row_number, columns["price"]).value
        if not port and raw_price in (None, ""):
            continue
        try:
            price = float(raw_price)
        except (TypeError, ValueError):
            skipped += 1
            continue
        if not grade or price <= 0:
            skipped += 1
            continue
        timestamp = _as_time(sheet.cell(row_number, columns.get("timestamp", 0)).value if columns.get("timestamp") else None, fallback)
        country = str(sheet.cell(row_number, columns.get("country", 0)).value or "") if columns.get("country") else ""
        region = str(sheet.cell(row_number, columns.get("region", 0)).value or "") if columns.get("region") else ""
        rows.append(_item(str(port), country, region, grade, price, timestamp, sheet.title))
    return rows, skipped


def _parse_wide(sheet, header_row, columns, fallback):
    rows = []
    skipped = 0
    timestamp_col = next((index for index, cell in enumerate(sheet[header_row], 1) if _text(cell.value) in HEADER_ALIASES["timestamp"]), None)
    country_col = next((index for index, cell in enumerate(sheet[header_row], 1) if _text(cell.value) in HEADER_ALIASES["country"]), None)
    for row_number in range(header_row + 1, min(sheet.max_row, 10000) + 1):
        port = sheet.cell(row_number, columns["port"]).value
        if not port:
            continue
        timestamp = _as_time(sheet.cell(row_number, timestamp_col).value if timestamp_col else None, fallback)
        country = str(sheet.cell(row_number, country_col).value or "") if country_col else ""
        for column, grade in columns["grade_columns"].items():
            raw_price = sheet.cell(row_number, column).value
            if raw_price in (None, ""):
                continue
            try:
                price = float(raw_price)
            except (TypeError, ValueError):
                skipped += 1
                continue
            if price <= 0:
                skipped += 1
                continue
            rows.append(_item(str(port), country, "", grade, price, timestamp, sheet.title))
    return rows, skipped


def _parse_matrix(sheet, fallback):
    best = None
    for grade_row in range(1, min(sheet.max_row, 12) + 1):
        grade_columns = {column: canonical_grade(sheet.cell(grade_row, column).value) for column in range(2, sheet.max_column + 1)}
        grade_columns = {column: grade for column, grade in grade_columns.items() if grade}
        if len(grade_columns) >= 2 and grade_row > 1:
            best = (grade_row, grade_columns)
            break
    if not best:
        return [], 0
    grade_row, grade_columns = best
    port_row = grade_row - 1
    ports = {}
    current_port = None
    for column in range(2, sheet.max_column + 1):
        value = sheet.cell(port_row, column).value
        if value not in (None, ""):
            current_port = str(value)
        ports[column] = current_port
    data_start = grade_row + 1
    for row_number in range(grade_row + 1, min(sheet.max_row, grade_row + 12) + 1):
        first = sheet.cell(row_number, 1).value
        if _looks_like_date(first):
            data_start = row_number
            break
    rows = []
    skipped = 0
    for row_number in range(data_start, min(sheet.max_row, 10000) + 1):
        row_time = _as_time(sheet.cell(row_number, 1).value, fallback)
        for column, grade in grade_columns.items():
            port = ports.get(column)
            raw_price = sheet.cell(row_number, column).value
            if not port or raw_price in (None, ""):
                continue
            try:
                price = float(raw_price)
            except (TypeError, ValueError):
                skipped += 1
                continue
            if price <= 0:
                skipped += 1
                continue
            rows.append(_item(port, "", "", grade, price, row_time, sheet.title))
    return rows, skipped


def _item(port, country, region, grade, price, timestamp, sheet_name):
    return ProviderObservation(
        port_name=port,
        country=country,
        region=region,
        grade=grade,
        price=price,
        currency="USD",
        unit="MT",
        source_time=timestamp,
        provenance_url="upload://excel",
        source_label=f"Excel · {sheet_name}",
    )


def parse_workbook(path: str | Path, file_name: str) -> ParseResult:
    fallback = utcnow()
    values_book = load_workbook(path, data_only=True, read_only=False)
    formulas_book = load_workbook(path, data_only=False, read_only=False)
    candidates = []
    total_missing = 0
    total_errors = 0
    for name in values_book.sheetnames:
        sheet = values_book[name]
        formula_sheet = formulas_book[name]
        missing, errors = _formula_diagnostics(sheet, formula_sheet)
        total_missing += missing
        total_errors += errors
        found = _find_columns(sheet)
        if found:
            header, columns, layout = found
            observations, skipped = _parse_long(sheet, header, columns, fallback) if layout == "long" else _parse_wide(sheet, header, columns, fallback)
        else:
            layout = "matrix"
            observations, skipped = _parse_matrix(sheet, fallback)
        if observations:
            candidates.append((len(observations), name, layout, observations, skipped))
    sheet_names = list(values_book.sheetnames)
    values_book.close()
    formulas_book.close()
    if not candidates:
        raise ValueError("No supported bunker-price table was detected in any worksheet")
    _, sheet_name, layout, observations, skipped = max(candidates, key=lambda item: item[0])
    deduped = {}
    for item in observations:
        identity = port_identity(item.port_name, item.country, item.region)
        deduped[(identity["code"], item.grade, item.source_time)] = item
    return ParseResult(
        file_name=file_name,
        sheet_name=sheet_name,
        layout=layout,
        observations=list(deduped.values()),
        skipped_cells=skipped,
        formula_cache_missing=total_missing,
        excel_errors=total_errors,
        sheet_names=sheet_names or [sheet_name],
    )


def import_excel_result(result: ParseResult) -> None:
    state = db.session.get(ProviderState, "excel")
    if not state:
        state = ProviderState(id="excel", name="Excel upload")
        db.session.add(state)
    upload = db.session.get(UploadState, 1)
    if not upload:
        upload = UploadState(id=1)
        db.session.add(upload)
    try:
        Observation.query.filter_by(provider_id="excel").delete(synchronize_session=False)
        port_cache = {port.code: port for port in Port.query.all()}
        for item in result.observations:
            identity = port_identity(item.port_name, item.country, item.region)
            port = port_cache.get(identity["code"])
            if not port:
                port = Port(**identity)
                db.session.add(port)
                port_cache[identity["code"]] = port
            db.session.add(Observation(
                provider_id="excel", port_code=identity["code"], grade=item.grade,
                price=item.price, currency="USD", unit="MT", source_time=item.source_time,
                retrieved_at=utcnow(), provenance_url=item.provenance_url, source_label=item.source_label,
            ))
        now = utcnow()
        state.configured = True
        state.status = "current"
        state.last_attempt_at = now
        state.last_success_at = now
        state.last_error = None
        state.records_last_run = len(result.observations)
        state.ports_last_run = len({port_identity(item.port_name)["code"] for item in result.observations})
        upload.file_name = result.file_name
        upload.sheet_name = result.sheet_name
        upload.uploaded_at = now
        upload.status = "current"
        upload.layout = result.layout
        upload.rows_received = len(result.observations)
        upload.ports_received = state.ports_last_run
        upload.skipped_cells = result.skipped_cells
        upload.formula_cache_missing = result.formula_cache_missing
        upload.excel_errors = result.excel_errors
        upload.last_error = None
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise


def mark_upload_error(message: str) -> None:
    upload = db.session.get(UploadState, 1)
    if not upload:
        upload = UploadState(id=1)
        db.session.add(upload)
    upload.status = "failed"
    upload.last_error = message[:1000]
    db.session.commit()
